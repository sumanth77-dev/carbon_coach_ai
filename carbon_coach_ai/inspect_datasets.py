"""
inspect_datasets.py - Analyze all XLSX files placed in datasets/
"""

import sys
from pathlib import Path
import openpyxl

PROJECT_ROOT = Path(__file__).resolve().parent
DATASETS_DIR = PROJECT_ROOT / "datasets"

def inspect_file(file_path: Path):
    print("=" * 70)
    print(f"FILE: {file_path.relative_to(PROJECT_ROOT)}")
    print(f"SIZE: {file_path.stat().st_size:,} bytes")
    
    try:
        # Load workbook with read_only for large files
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        print(f"SHEETS ({len(wb.sheetnames)}): {wb.sheetnames[:10]}...")
        
        for name in wb.sheetnames[:4]:
            sheet = wb[name]
            print(f"\n--- Sheet: '{name}' ---")
            rows = []
            for i, row in enumerate(sheet.iter_rows(values_only=True)):
                if any(row):
                    rows.append(row[:8])
                if len(rows) >= 6:
                    break
            for r in rows:
                print("  ", [str(c)[:30] if c is not None else "" for c in r])
        wb.close()
    except Exception as e:
        print(f"ERROR reading {file_path.name}: {e}")

def main():
    xlsx_files = list(DATASETS_DIR.glob("**/*.xlsx"))
    print(f"Total XLSX files found: {len(xlsx_files)}")
    for f in sorted(xlsx_files):
        inspect_file(f)

if __name__ == "__main__":
    main()
